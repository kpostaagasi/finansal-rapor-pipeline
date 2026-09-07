#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TEFAS net fon akışı raporu üreticisi.

Net akış = (o günün tedavüldeki pay sayısı − önceki işlem gününün pay sayısı)
           × o günün pay fiyatı   [fon bazında hesaplanıp toplanır]

Evren (her iki taraf da ALTIN fonları):
  YF  — adında ALTIN geçen yatırım fonları ("ALTINCI"/"ON ALTINCI" elenir,
        "GOLD" geçenler eklenir, "GOLDEN ..." elenir)
  EYF — `eyf_fonlar.json`daki altın emeklilik fonları (kullanıcının TEFAS'tan
        dışa aktardığı Excel listesi; liste değişirse o dosya güncellenip
        --bootstrap ile seri yeniden kurulmalı)

Veri `akis_veri.json`'da birikir; her çalışmada yalnızca son ~12 gün yeniden
çekilip üstüne yazılır (TEFAS pencere sınırı 28 gün, rate limit ~6 istek/dk).
Önbellek yoksa BAS_TARIH'ten bugüne tam kurulum yapılır (~8 dk sürer).

Kullanım:
  python3 tefas_akis.py             # artımlı güncelle + HTML üret
  python3 tefas_akis.py --bootstrap # önbelleği sıfırdan kur
  python3 tefas_akis.py --no-fetch  # veri çekme, mevcut önbellekten HTML üret

Çıktı: tefas_net_akis.html — "OK: <gün sayısı> gün, <yf> altın / <eyf> emeklilik fonu"
"""
import os, sys, json, time, tempfile, datetime as dt
from collections import defaultdict

import requests

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "akis_veri.json")
FON_CACHE = os.path.join(HERE, "fon_veri.json")   # fon bazında fiyat/pay (Excel için)
EXCEL_OUT = os.path.expanduser("~/Documents/TEFAS_Altin_Fonlari_Akis.xlsx")
EYF_LISTE = os.path.join(HERE, "eyf_fonlar.json")
TEMPLATE = os.path.join(HERE, "tefas_template.html")
OUT = os.path.join(HERE, "tefas_net_akis.html")
DESKTOP_COPY = os.path.expanduser("~/Documents/TEFAS_Net_Akis_Grafik.html")

BAS_TARIH = dt.date(2024, 12, 26)   # serinin başlangıcı (ilk gün akış için referans, seriye girmez)
PENCERE = 25                        # gün; TEFAS sınırı 28
INCREMENTAL_GUN = 12                # her gün yeniden hesaplanan kuyruk
YAKIN_PENCERE_GUN = 90              # "son 90 gün" penceresi: panonun gap_periods varsayılanıyla aynı

API = "https://www.tefas.gov.tr/api/funds/fonGnlBlgSiraliGetir"
HEADERS = {
    "Content-Type": "application/json",
    "Origin": "https://www.tefas.gov.tr",
    "Referer": "https://www.tefas.gov.tr/",
    "User-Agent": ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"),
}


def log(msg):
    print(f"{dt.datetime.now():%H:%M:%S} — {msg}", file=sys.stderr, flush=True)


def script_json(value):
    """JSON'u HTML <script> bağlamında güvenli biçimde kodlar."""
    return (json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            .replace("<", "\\u003c")
            .replace(">", "\\u003e")
            .replace("&", "\\u0026")
            .replace("\u2028", "\\u2028")
            .replace("\u2029", "\\u2029"))


def atomic_json_dump(path, value):
    """Tam JSON'u aynı dizinde yazıp fsync sonrası atomik olarak yer değiştirir."""
    directory = os.path.dirname(os.path.abspath(path))
    os.makedirs(directory, exist_ok=True)
    fd, temporary = tempfile.mkstemp(
        prefix=f".{os.path.basename(path)}.", suffix=".tmp", dir=directory
    )
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(value, f, ensure_ascii=False)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def altin_fonu(unvan):
    """Adı altın fonuna işaret ediyor mu? (sıra sayısı 'altıncı' tuzağına dikkat)"""
    u = unvan.upper()
    u_temiz = u.replace("ON ALTINCI", " ").replace("ONALTINCI", " ").replace("ALTINCI", " ")
    if "ALTIN" in u_temiz:
        return True
    return "GOLD" in u and "GOLDEN" not in u


def eyf_kodlari():
    """Rapora girecek altın emeklilik fonlarının kodları (Excel'den türetilmiş liste)."""
    with open(EYF_LISTE, encoding="utf-8") as f:
        return set(json.load(f)["fonlar"])


def fetch(fon_tipi, bas, bit, deneme=3):
    body = {"fonTipi": fon_tipi, "fonKodu": None, "aramaMetni": None, "fonTurKod": None,
            "fonGrubu": None, "sfonTurKod": None,
            "basTarih": bas.strftime("%Y%m%d"), "bitTarih": bit.strftime("%Y%m%d"),
            "basSira": 1, "bitSira": 100000, "fonTurAciklama": None,
            "dil": "TR", "kurucuKod": None}
    for i in range(1, deneme + 1):
        try:
            r = requests.post(API, headers=HEADERS, json=body, timeout=120)
            r.raise_for_status()
            return r.json()["resultList"] or []
        except Exception as e:
            log(f"  {fon_tipi} {bas}–{bit} deneme {i}/{deneme} hata: {str(e)[:120]}")
            if i == deneme:
                raise
            time.sleep(20)


def gecerli_gozlem(kayit):
    """TEFAS kaydı geçerli bir gözlem mi?

    Bozuk bir fetch bazen tedavüldeki pay sayısını (tedPaySayisi) ve fiyatı
    0 döndürüyor; pay×fiyat=0 matematiksel olarak geçerli bir 0 akış gibi
    görünüp fail-closed kontrollerini atlatıyor, "Hesaplandı" etiketiyle 0 TL
    yayınlanıyordu. tedPaySayisi<=0 veya fiyat<=0 olan kayıt gözlem SAYILMAZ
    — veri boşluğu olarak ele alınır.
    """
    pay, fiyat = kayit
    return pay is not None and fiyat is not None and pay > 0 and fiyat > 0


BOLUNME_PAY_ESIGI = 1.5         # pay oranı bu eşiği aşarsa/altına inerse bölünme adayı
BOLUNME_DEGER_TOLERANSI = 0.05  # pay×fiyat oranı 1'den bu kadar sapabilir


def pay_bolunmesi(onceki_kayit, simdi_kayit):
    """İki ardışık geçerli gözlem arasında pay bölünmesi/birleşmesi var mı?

    TEFAS bazen bir fonun pay bölünmesini/birleşmesini sıradan bir güncelleme
    gibi veriyor: pay sayısı binlerce kat sıçrıyor, fiyat ters orantılı düşüyor
    (ör. TI2 2025-01-20: pay ×9971, fiyat ÷9834). (pay_t - pay_onceki) × fiyat_t
    formülü bunu matematiksel olarak geçerli ama hayali, milyarlarca TL'lik bir
    akış gibi hesaplıyor. Pay oranı BOLUNME_PAY_ESIGI'yi aşıp/altına inip fon
    değeri (pay×fiyat) ~sabit kalıyorsa (BOLUNME_DEGER_TOLERANSI içinde) bu
    sıradan bir alım/satım değil, birim değişimidir — akış hesaplanamaz
    (fail-closed: bkz. akis_serisi).
    """
    onceki_pay, onceki_fiyat = onceki_kayit
    simdi_pay, simdi_fiyat = simdi_kayit
    pay_orani = simdi_pay / onceki_pay
    fiyat_orani = simdi_fiyat / onceki_fiyat
    deger_orani = pay_orani * fiyat_orani
    esik_asildi = pay_orani >= BOLUNME_PAY_ESIGI or pay_orani <= 1 / BOLUNME_PAY_ESIGI
    return esik_asildi and abs(deger_orani - 1) < BOLUNME_DEGER_TOLERANSI


def topla(bas, bit, adlar=None):
    """[bas, bit] aralığı için fon bazında {kod: {tarih: (pay, fiyat)}} döndürür."""
    yf, eyf = defaultdict(dict), defaultdict(dict)
    eyf_kod = eyf_kodlari()
    pencere_bas = bas
    while pencere_bas <= bit:
        pencere_bit = min(pencere_bas + dt.timedelta(days=PENCERE), bit)
        for tip, hedef in (("YAT", yf), ("EMK", eyf)):
            for x in fetch(tip, pencere_bas, pencere_bit):
                if tip == "YAT" and not altin_fonu(x["fonUnvan"]):
                    continue
                if tip == "EMK" and x["fonKodu"] not in eyf_kod:
                    continue
                if not gecerli_gozlem((x["tedPaySayisi"], x["fiyat"])):
                    continue
                hedef[x["fonKodu"]][x["tarih"]] = (x["tedPaySayisi"], x["fiyat"])
                if adlar is not None:
                    adlar[x["fonKodu"]] = x["fonUnvan"]
            time.sleep(10)   # rate limit ~6 istek/dk
        log(f"  çekildi {pencere_bas} – {pencere_bit}")
        pencere_bas = pencere_bit + dt.timedelta(days=1)
    return yf, eyf


def fon_verisi_kaydet(yf_fon, eyf_fon, adlar):
    """Fon bazında fiyat/pay verisini biriktirir (Excel çıktısı buradan üretilir)."""
    veri = {"yf": {}, "eyf": {}, "ad": {}}
    if os.path.exists(FON_CACHE):
        with open(FON_CACHE, encoding="utf-8") as f:
            veri = json.load(f)
    for anahtar, kaynak in (("yf", yf_fon), ("eyf", eyf_fon)):
        hedef = veri.setdefault(anahtar, {})
        for kod, seri in kaynak.items():
            g = hedef.setdefault(kod, {})
            for tarih, (pay, fiyat) in seri.items():
                g[tarih] = [pay, fiyat]
    veri.setdefault("ad", {}).update(adlar)
    atomic_json_dump(FON_CACHE, veri)
    return veri


def akislari_hesapla(fonlar, tarihler):
    """Akışları yalnız evrenin ardışık TEFAS veri tarihleri arasında toplar.

    Önceki veya güncel tarihte eksik fon varsa kısmi toplam doğruymuş gibi
    gösterilmez; o tarih için toplam ``None`` olur.

    Bozuk TEFAS kaydı (bkz. gecerli_gozlem — tedPaySayisi<=0 veya fiyat<=0)
    gözlem sayılmaz: seri bu fonksiyonda okunurken elenir, o fon o tarihte
    hiç gözlem vermemiş gibi ele alınır (veri boşluğu).

    Pay bölünmesi/birleşmesi tespit edilen fon-gün (bkz. pay_bolunmesi) de
    aynı şekilde toplamdan dışlanır ve günün toplamını None'a çeker: tek bir
    üyenin bölünmesi grup toplamını şişirmemeli (fail-closed).
    """
    toplam: dict[str, float | None] = {t: 0.0 for t in tarihler}
    sayim = {t: 0 for t in tarihler}
    sorun = {t: {"missing_current": [], "missing_previous": []} for t in tarihler}
    gecerli = {
        kod: {t: v for t, v in seri.items() if gecerli_gozlem(v)}
        for kod, seri in fonlar.items()
    }
    for t in tarihler:
        sayim[t] = sum(t in seri for seri in gecerli.values())
    for onceki, t in zip(tarihler, tarihler[1:]):
        onceki_kodlar = {kod for kod, seri in gecerli.items() if onceki in seri}
        guncel_kodlar = {kod for kod, seri in gecerli.items() if t in seri}
        sorun[t]["missing_current"] = sorted(onceki_kodlar - guncel_kodlar)
        sorun[t]["missing_previous"] = sorted(guncel_kodlar - onceki_kodlar)
        ortak_kodlar = onceki_kodlar & guncel_kodlar
        bolunenler = [kod for kod in ortak_kodlar
                      if pay_bolunmesi(gecerli[kod][onceki], gecerli[kod][t])]
        for kod in ortak_kodlar - set(bolunenler):
            pay, fiyat = gecerli[kod][t]
            toplam[t] = (toplam[t] or 0.0) + (pay - gecerli[kod][onceki][0]) * fiyat
        # İlk kez gözlemlenen fon (piyasaya çıkış) akış boşluğu DEĞİLDİR; akışı
        # hesaplanamaz ama o günün toplamını iptal etmez. metadata'da görünmeye
        # devam eder; html_uret launch_only ile recent_gap'e saymaz. HAM (bkz.
        # fonlar, filtrelenmemiş) min tarih kullanılır: fonun daha eski bozuk
        # bir kaydı varsa bu gerçek bir çıkış değildir, kopukluk sayılır.
        kopuklar = [kod for kod in (guncel_kodlar - onceki_kodlar)
                    if min(fonlar[kod]) != t]
        if sorun[t]["missing_current"] or kopuklar or bolunenler:
            toplam[t] = None
    return toplam, sayim, sorun


KESINLESME_SAATI = 16  # TEFAS netleşmesi ~15:27; bu saatten önce bugünün satırı öncül kabul edilir


def kesin_tarihler(tarihler):
    """Netleşme saatinden önce bugünün (öncül) satırını seriden düşer."""
    simdi = dt.datetime.now()
    if simdi.hour < KESINLESME_SAATI and tarihler and tarihler[-1] == dt.date.today().isoformat():
        log(f"bugünün öncül satırı düşüldü ({dt.date.today()} — netleşme {KESINLESME_SAATI}:00 öncesi)")
        return tarihler[:-1]
    return tarihler


def grup_onbellegi_uret(yf_fonlar, eyf_fonlar, guncelleme=None):
    """Tam fon cache'inden dürüst grup serisini ve kapsam metadata'sını kurar."""
    tarihler = sorted({t for s in yf_fonlar.values() for t in s} |
                      {t for s in eyf_fonlar.values() for t in s})
    if len(tarihler) < 2:
        raise RuntimeError("grup akışı için en az iki TEFAS veri tarihi gerekli")
    yf_akis, yf_say, yf_sorun = akislari_hesapla(yf_fonlar, tarihler)
    eyf_akis, eyf_say, eyf_sorun = akislari_hesapla(eyf_fonlar, tarihler)
    gunler = kesin_tarihler(tarihler[1:])

    def yuvarla(value):
        return round(value) if value is not None else None

    onbellek = {
        "d": gunler,
        "yf": [yuvarla(yf_akis[t]) for t in gunler],
        "eyf": [yuvarla(eyf_akis[t]) for t in gunler],
        "yf_ilk": {kod: min(seri) for kod, seri in yf_fonlar.items()},
        "eyf_ilk": {kod: min(seri) for kod, seri in eyf_fonlar.items()},
        "yf_n": {t: yf_say[t] for t in gunler},
        "eyf_n": {t: eyf_say[t] for t in gunler},
        "guncelleme": guncelleme or dt.datetime.now().isoformat(timespec="seconds"),
    }
    for anahtar, say, sorunlar in (
        ("yf", yf_say, yf_sorun), ("eyf", eyf_say, eyf_sorun)
    ):
        onbellek[f"{anahtar}_expected_n"] = {
            t: say[t] + len(sorunlar[t]["missing_current"]) for t in gunler
        }
        onbellek[f"{anahtar}_calc_n"] = {
            t: say[t] - len(sorunlar[t]["missing_previous"]) for t in gunler
        }
        onbellek[f"{anahtar}_missing"] = {
            t: sorunlar[t]["missing_current"] for t in gunler
        }
        onbellek[f"{anahtar}_uncomputed"] = {
            t: sorunlar[t]["missing_previous"] for t in gunler
        }
    return onbellek


def veri_guncelle(tam=False):
    onbellek = {
        "d": [], "yf": [], "eyf": [], "yf_n": {}, "eyf_n": {},
        "yf_expected_n": {}, "eyf_expected_n": {},
        "yf_calc_n": {}, "eyf_calc_n": {},
        "yf_missing": {}, "eyf_missing": {},
        "yf_uncomputed": {}, "eyf_uncomputed": {},
    }
    if not tam and os.path.exists(CACHE):
        with open(CACHE, encoding="utf-8") as f:
            onbellek = json.load(f)

    bugun = dt.date.today()
    if onbellek["d"]:
        son = dt.date.fromisoformat(onbellek["d"][-1])
        bas = max(BAS_TARIH, son - dt.timedelta(days=INCREMENTAL_GUN))
    else:
        bas = BAS_TARIH
    log(f"veri çekiliyor: {bas} → {bugun}")

    adlar = {}
    yf_fon, eyf_fon = topla(bas, bugun, adlar)
    cekilen_tarihler = sorted({t for s in yf_fon.values() for t in s} |
                              {t for s in eyf_fon.values() for t in s})
    if not cekilen_tarihler:
        raise RuntimeError("TEFAS'tan veri gelmedi")
    birlesik = fon_verisi_kaydet(yf_fon, eyf_fon, adlar)
    onbellek = grup_onbellegi_uret(
        birlesik.get("yf", {}), birlesik.get("eyf", {}),
        dt.datetime.now().isoformat(timespec="seconds"),
    )
    atomic_json_dump(CACHE, onbellek)
    return onbellek


def html_uret(onbellek):
    with open(TEMPLATE, encoding="utf-8") as f:
        html = f.read()
    d = onbellek["d"]
    raw = script_json({"d": d, "yf": onbellek["yf"], "eyf": onbellek["eyf"]})
    say = lambda tablo, tarih, varsayilan: str(tablo.get(tarih, varsayilan))
    ilk, son = d[0], d[-1]
    yf_son = onbellek.get("yf_n", {}).get(son)
    eyf_son = onbellek.get("eyf_n", {}).get(son)
    yf_beklenen = onbellek.get("yf_expected_n", {}).get(son, yf_son)
    eyf_beklenen = onbellek.get("eyf_expected_n", {}).get(son, eyf_son)
    bulunan = (yf_son + eyf_son) if isinstance(yf_son, int) and isinstance(eyf_son, int) else None
    beklenen = (yf_beklenen + eyf_beklenen
                if isinstance(yf_beklenen, int) and isinstance(eyf_beklenen, int) else None)
    eksik = (onbellek.get("yf_missing", {}).get(son, [])
             + onbellek.get("eyf_missing", {}).get(son, []))
    hesaplanamayan = (onbellek.get("yf_uncomputed", {}).get(son, [])
                      + onbellek.get("eyf_uncomputed", {}).get(son, []))
    yf_hesaplanan = onbellek.get("yf_calc_n", {}).get(son, yf_son)
    eyf_hesaplanan = onbellek.get("eyf_calc_n", {}).get(son, eyf_son)
    hesaplanan = (yf_hesaplanan + eyf_hesaplanan
                  if isinstance(yf_hesaplanan, int) and isinstance(eyf_hesaplanan, int) else None)
    # Yakın pencere: son veri gününden YAKIN_PENCERE_GUN (90) gün geriye — dahil.
    kesim = dt.date.fromisoformat(son) - dt.timedelta(days=YAKIN_PENCERE_GUN)
    yakin_eksik_gun = 0
    ilk_yf = onbellek.get("yf_ilk", {})
    ilk_eyf = onbellek.get("eyf_ilk", {})
    for tarih, yf, eyf in zip(d, onbellek["yf"], onbellek["eyf"]):
        if dt.date.fromisoformat(tarih) < kesim:
            continue
        if yf is not None and eyf is not None:
            continue
        yf_mis = onbellek.get("yf_missing", {}).get(tarih, [])
        eyf_mis = onbellek.get("eyf_missing", {}).get(tarih, [])
        yf_un = onbellek.get("yf_uncomputed", {}).get(tarih, [])
        eyf_un = onbellek.get("eyf_uncomputed", {}).get(tarih, [])
        launch_only = (
            not yf_mis
            and not eyf_mis
            and all(ilk_yf.get(k) == tarih for k in yf_un)
            and all(ilk_eyf.get(k) == tarih for k in eyf_un)
        )
        if not launch_only:
            yakin_eksik_gun += 1
    report_meta = {
        "last_successful_run": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "data_end_date": son,
        "expected_count": beklenen,
        "found_count": bulunan,
        "missing": eksik,
        "uncomputed_count": ((beklenen - hesaplanan)
                             if isinstance(beklenen, int) and isinstance(hesaplanan, int) else None),
        "uncomputed": hesaplanamayan,
        "recent_gap_count": yakin_eksik_gun,
        "count_label": "fon",
        "source": "TEFAS",
    }
    g, a, y = ilk.split("-")[2], ilk.split("-")[1], ilk.split("-")[0]
    html = (html.replace("__RAW__", raw)
                .replace("__REPORT_META__", script_json(report_meta))
                .replace("__YF_N__", say(onbellek.get("yf_n", {}), son, "?"))
                .replace("__EYF_N__", say(onbellek.get("eyf_n", {}), son, "?"))
                .replace("__YF0_N__", say(onbellek.get("yf_n", {}), ilk, "?"))
                .replace("__EYF0_N__", say(onbellek.get("eyf_n", {}), ilk, "?"))
                .replace("__ILK_TARIH__", f"{g}.{a}.{y}"))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(html)
    try:
        with open(DESKTOP_COPY, "w", encoding="utf-8") as f:
            f.write(html)
    except OSError as e:
        log(f"masaüstü kopyası yazılamadı: {e}")
    return html


def excel_uret(yol=EXCEL_OUT):
    """Fiyat / Tedavül / Net Akış sayfalarından oluşan çalışma kitabı üretir.

    Net Akış = (o günün tedavüldeki pay sayısı − önceki işlem gününün pay sayısı)
               × o günün fiyatı   [yani tedavül değişimi, DEĞİŞİMİN OLDUĞU günün fiyatıyla]
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    with open(FON_CACHE, encoding="utf-8") as f:
        veri = json.load(f)
    adlar = veri.get("ad", {})

    gruplar = [("yf", "Altın Yatırım Fonu"), ("eyf", "Altın Emeklilik Fonu")]
    kodlar, grup_ad = [], {}
    for anahtar, etiket in gruplar:
        for kod in sorted(veri.get(anahtar, {})):
            kodlar.append((anahtar, kod))
            grup_ad[kod] = etiket
    tarihler = kesin_tarihler(sorted({t for a, _ in gruplar for s in veri.get(a, {}).values() for t in s}))

    wb = Workbook()
    wb.remove(wb.active)
    kalin = Font(bold=True)

    def sayfa(baslik, deger_fn, bicim):
        ws = wb.create_sheet(baslik)
        ws.cell(1, 1, "Fon Kodu").font = kalin
        ws.cell(2, 1, "Fon Adı").font = kalin
        ws.cell(3, 1, "Tür").font = kalin
        ws.cell(4, 1, "Tarih").font = kalin
        for j, (anahtar, kod) in enumerate(kodlar, start=2):
            ws.cell(1, j, kod).font = kalin
            ws.cell(2, j, adlar.get(kod, "")).alignment = Alignment(wrap_text=False)
            ws.cell(3, j, grup_ad[kod])
        for i, t in enumerate(tarihler, start=5):
            ws.cell(i, 1, dt.date.fromisoformat(t)).number_format = "DD.MM.YYYY"
            for j, (anahtar, kod) in enumerate(kodlar, start=2):
                v = deger_fn(veri[anahtar][kod], t)
                if v is not None:
                    c = ws.cell(i, j, v)
                    c.number_format = bicim
        ws.freeze_panes = "B5"
        ws.column_dimensions["A"].width = 12
        for j in range(2, len(kodlar) + 2):
            ws.column_dimensions[get_column_letter(j)].width = 16
        return ws

    sayfa("Fiyat", lambda s, t: s[t][1] if t in s else None, "0.000000")
    sayfa("Tedavüldeki Pay Sayısı", lambda s, t: s[t][0] if t in s else None, "#,##0.00")

    onceki_rapor = {t: tarihler[i - 1] for i, t in enumerate(tarihler) if i}

    def akis(s, t):
        o = onceki_rapor.get(t)
        if o is None or t not in s:
            return None
        if o not in s:
            return None
        return (s[t][0] - s[o][0]) * s[t][1]

    sayfa("Net Akış", akis, "#,##0")

    # Özet: günlük grup toplamları (grafikteki seriyle aynı; kısmi günler boş)
    grup_toplamlari = {
        anahtar: akislari_hesapla(veri[anahtar], tarihler)[0]
        for anahtar, _ in gruplar
    }
    ws = wb.create_sheet("Özet", 0)
    for j, b in enumerate(["Tarih", "Altın Yatırım Fonları (TL)",
                           "Altın Emeklilik Fonları (TL)", "Toplam (TL)"], start=1):
        ws.cell(1, j, b).font = kalin
    # ilk gün yalnızca referans (önceki günü yok) — Özet'e girmez
    for i, t in enumerate(tarihler[1:], start=2):
        toplamlar = [grup_toplamlari[anahtar][t] for anahtar, _ in gruplar]
        ws.cell(i, 1, dt.date.fromisoformat(t)).number_format = "DD.MM.YYYY"
        for j, v in enumerate(toplamlar, start=2):
            if v is not None:
                ws.cell(i, j, round(v)).number_format = "#,##0"
        tam_toplamlar = [v for v in toplamlar if v is not None]
        if len(tam_toplamlar) == len(toplamlar):
            ws.cell(i, 4, round(sum(tam_toplamlar))).number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 12
    for h in "BCD":
        ws.column_dimensions[h].width = 26

    wb.save(yol)
    return yol, len(tarihler), len(kodlar)


def main():
    no_fetch = "--no-fetch" in sys.argv
    tam = "--bootstrap" in sys.argv
    if no_fetch:
        if os.path.exists(FON_CACHE):
            with open(FON_CACHE, encoding="utf-8") as f:
                fon_veri = json.load(f)
            onbellek = grup_onbellegi_uret(
                fon_veri.get("yf", {}), fon_veri.get("eyf", {}),
                dt.datetime.now().isoformat(timespec="seconds"),
            )
            atomic_json_dump(CACHE, onbellek)
        else:
            with open(CACHE, encoding="utf-8") as f:
                onbellek = json.load(f)
    else:
        onbellek = veri_guncelle(tam=tam)
    html_uret(onbellek)
    if os.path.exists(FON_CACHE):
        try:
            yol, gun, fon = excel_uret()
            log(f"Excel: {yol} ({gun} gün × {fon} fon)")
        except Exception as e:
            log(f"Excel üretilemedi: {e}")
    son = onbellek["d"][-1]
    print(f"OK: {len(onbellek['d'])} gün ({onbellek['d'][0]} → {son}) — "
          f"{onbellek.get('yf_n', {}).get(son, '?')} altın / "
          f"{onbellek.get('eyf_n', {}).get(son, '?')} emeklilik fonu")
    meta_eksik = (onbellek.get("yf_missing", {}).get(son, [])
                  + onbellek.get("eyf_missing", {}).get(son, []))
    meta_hesaplanamayan = (onbellek.get("yf_uncomputed", {}).get(son, [])
                           + onbellek.get("eyf_uncomputed", {}).get(son, []))
    if onbellek["yf"][-1] is None or onbellek["eyf"][-1] is None:
        log("HATA: son TEFAS tarihinde grup akışı tam hesaplanamadı; dağıtım engellendi")
        if meta_eksik:
            log("  son tarihte eksik fonlar: " + ", ".join(meta_eksik))
        if meta_hesaplanamayan:
            log("  önceki gözlemi olmayan fonlar: " + ", ".join(meta_hesaplanamayan))
        return 4
    return 0


if __name__ == "__main__":
    sys.exit(main())
